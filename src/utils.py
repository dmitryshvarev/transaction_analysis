import json
import logging

from datetime import datetime
from typing import Dict, List, Union

from openpyxl import load_workbook

from src.external_api import get_exchange_rate

logger = logging.getLogger(__name__)
file_handler = logging.FileHandler(r"..\logs\utils.log", mode="w", encoding="utf-8")
file_formatter = logging.Formatter("%(asctime)s %(filename)s %(levelname)s: %(message)s")
file_handler.setFormatter(file_formatter)
logger.addHandler(file_handler)
logger.setLevel(logging.DEBUG)


def get_transactions_for_analysis(file_path: str, user_date: str) -> List[Dict[str, Union[str, float, int]]]:
    """Получаем данные о финансовых операциях с начала месяца по user_date"""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    headers = next(rows)

    transactions = []

    user_datetime: datetime = datetime.strptime(user_date, "%Y-%m-%d %H:%M:%S")
    start_of_month = datetime(user_datetime.year, user_datetime.month, 1)

    for row in rows:
        row_dict = dict(zip(headers, row))
        trans_date = row_dict.get("Дата операции")
        if trans_date:
            trans_date = datetime.strptime(trans_date, "%d.%m.%Y %H:%M:%S")

        # Фильтруем записи: от начала месяца до пользовательской даты включительно
        if start_of_month <= trans_date <= user_datetime:
            transactions.append(row_dict)

    wb.close()
    if transactions:
        logger.info("Список транзакций сформирован")
        return transactions
    else:
        logger.warning("Список транзакций пуст")


def amount_in_rub(transaction: Dict[str, Union[str, float, int]]) -> float:
    """Возвращает сумму транзакции в рублях"""
    amount = 0
    currency_code = transaction.get("Валюта платежа")
    if currency_code:
        logger.info("Получаем данные о сумме транзакции")
        if currency_code == "RUB":
            logger.info("Данная транзакция в рублях")
            amount = float(transaction.get("Сумма платежа"))
        else:
            logger.info("Данная транзакция не в рублях. Происходит обращение к API")
            exchange_rate = get_exchange_rate()["Valute"][currency_code]["Value"]
            amount = round(float(transaction.get("Сумма платежа")) * exchange_rate, 2)
    else:
        logger.error("Неверные данные о транзакции")

    return amount


def get_settings(settings_path: str) -> Dict[str, list]:
    """Получаем данные из файла настроек"""
    settings_data = []
    try:
        with open(settings_path, "r", encoding="utf-8") as settings_file:
            logger.info(f'Файл "{settings_path}" успешно загружен')
            try:
                settings_data = json.load(settings_file)
                logger.info(f'Данные из файла "{settings_path}" преобразованы в объект Python')
            except json.JSONDecodeError:
                logger.error(f'Ошибка декодирования файла "{settings_path}"')
                return settings_data
    except FileNotFoundError:
        logger.error(f'Файл "{settings_path}" не найден')
        return settings_data

    return settings_data
